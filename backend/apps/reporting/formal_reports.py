import csv
from datetime import date
from io import BytesIO, StringIO
from uuid import UUID

from django.db.models import Q
from rest_framework.exceptions import ValidationError

from apps.actions.models import Action
from apps.assessments.models import Assessment, AssessmentItem
from apps.findings.models import Finding
from apps.identity.services import accessible_organization_unit_ids, has_whole_tenant_permission
from apps.internal_audits.models import AuditEngagement, Workpaper
from apps.organizations.models import OrganizationUnit
from apps.risks.models import RiskTreatment


REPORT_CATALOG = {
    "soa": {
        "title": "بیانیه کاربردپذیری (SoA)",
        "formats": ["pdf", "docx", "xlsx", "csv"],
        "filters": ["organization_unit", "assessment"],
    },
    "rtp": {
        "title": "برنامه برخورد با ریسک (RTP)",
        "formats": ["pdf", "docx", "xlsx", "csv"],
        "filters": ["organization_unit", "from_date", "to_date", "status"],
    },
    "audit": {
        "title": "گزارش ممیزی",
        "formats": ["pdf", "docx"],
        "filters": ["organization_unit", "engagement"],
    },
    "findings": {
        "title": "دفتر یافته‌ها",
        "formats": ["pdf", "xlsx", "csv"],
        "filters": ["organization_unit", "from_date", "to_date", "status", "severity"],
    },
    "actions": {
        "title": "دفتر اقدامات",
        "formats": ["pdf", "xlsx", "csv"],
        "filters": ["organization_unit", "from_date", "to_date", "status", "priority"],
    },
    "executive": {
        "title": "خلاصه مدیریتی",
        "formats": ["pdf", "docx", "xlsx"],
        "filters": ["organization_unit"],
    },
}


def _parse_uuid(value, field):
    if not value:
        return None
    try:
        return UUID(str(value))
    except (TypeError, ValueError, AttributeError) as exc:
        raise ValidationError({field: "Invalid UUID."}) from exc


def _parse_date(value, field):
    if not value:
        return None
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError) as exc:
        raise ValidationError({field: "Use YYYY-MM-DD."}) from exc


def normalized_filters(request, tenant, permission="report.view"):
    whole = has_whole_tenant_permission(request.user, tenant, permission)
    allowed = set(accessible_organization_unit_ids(request.user, tenant, permission))
    requested_unit = _parse_uuid(request.query_params.get("organization_unit"), "organization_unit")
    if requested_unit:
        exists = OrganizationUnit.objects.for_tenant(tenant).filter(
            id=requested_unit, deleted_at__isnull=True
        ).exists()
        if not exists:
            raise ValidationError({"organization_unit": "Organization unit not found."})
        if not whole and requested_unit not in allowed:
            raise ValidationError({"organization_unit": "Organization unit is outside your report scope."})
        scoped_units = {requested_unit}
    else:
        scoped_units = None if whole else allowed

    from_date = _parse_date(request.query_params.get("from_date"), "from_date")
    to_date = _parse_date(request.query_params.get("to_date"), "to_date")
    if from_date and to_date and from_date > to_date:
        raise ValidationError({"date_range": "from_date must be on or before to_date."})

    return {
        "organization_unit": requested_unit,
        "allowed_units": scoped_units,
        "from_date": from_date,
        "to_date": to_date,
        "status": request.query_params.get("status") or "",
        "severity": request.query_params.get("severity") or "",
        "priority": request.query_params.get("priority") or "",
        "assessment": _parse_uuid(request.query_params.get("assessment"), "assessment"),
        "engagement": _parse_uuid(request.query_params.get("engagement"), "engagement"),
    }


def public_filters(filters):
    return {
        key: (str(value) if value is not None else None)
        for key, value in filters.items()
        if key != "allowed_units"
    }


def _scope_queryset(queryset, filters):
    units = filters["allowed_units"]
    if units is not None:
        queryset = queryset.filter(organization_unit_id__in=units)
    return queryset


def _display_user(user):
    if not user:
        return ""
    return user.get_full_name().strip() or user.get_username()


def build_soa(tenant, filters):
    if not filters["assessment"]:
        raise ValidationError({"assessment": "Assessment is required for SoA."})
    assessments = _scope_queryset(
        Assessment.objects.filter(tenant=tenant, deleted_at__isnull=True), filters
    )
    assessment = assessments.filter(id=filters["assessment"]).select_related(
        "framework_version__framework", "organization_unit"
    ).first()
    if not assessment:
        raise ValidationError({"assessment": "Assessment not found in your report scope."})

    rows = []
    items = AssessmentItem.objects.filter(
        assessment=assessment, deleted_at__isnull=True
    ).order_by("requirement__sort_order", "requirement_code_snapshot")
    for item in items:
        rows.append([
            item.requirement_code_snapshot,
            item.requirement_title_snapshot,
            item.get_applicability_display(),
            item.get_status_display(),
            "" if item.score is None else str(item.score),
            item.assessor_comment,
            item.reviewer_comment,
        ])
    return {
        "title": f"SoA — {assessment.title}",
        "subtitle": f"{assessment.framework_version.framework.name} / {assessment.framework_version.version_code}",
        "columns": ["Requirement", "Title", "Applicability", "Status", "Score", "Assessor comment", "Reviewer comment"],
        "rows": rows,
        "summary": {"assessment_id": str(assessment.id), "row_count": len(rows)},
    }


def build_rtp(tenant, filters):
    qs = RiskTreatment.objects.filter(
        risk__tenant=tenant, deleted_at__isnull=True, risk__deleted_at__isnull=True
    ).select_related("risk", "risk__organization_unit", "owner")
    units = filters["allowed_units"]
    if units is not None:
        qs = qs.filter(risk__organization_unit_id__in=units)
    if filters["status"]:
        qs = qs.filter(status=filters["status"])
    if filters["from_date"]:
        qs = qs.filter(target_date__gte=filters["from_date"])
    if filters["to_date"]:
        qs = qs.filter(target_date__lte=filters["to_date"])
    rows = [[
        x.risk.code,
        x.risk.title,
        x.get_strategy_display(),
        x.description,
        _display_user(x.owner),
        x.target_date.isoformat() if x.target_date else "",
        "" if x.target_score is None else str(x.target_score),
        x.get_approval_status_display(),
        x.get_status_display(),
    ] for x in qs.order_by("target_date", "risk__code")]
    return {
        "title": "Risk Treatment Plan",
        "columns": ["Risk", "Title", "Strategy", "Treatment", "Owner", "Target date", "Target score", "Approval", "Status"],
        "rows": rows,
        "summary": {"row_count": len(rows)},
    }


def build_audit(tenant, filters):
    if not filters["engagement"]:
        raise ValidationError({"engagement": "Audit engagement is required."})
    qs = _scope_queryset(
        AuditEngagement.objects.filter(tenant=tenant, deleted_at__isnull=True), filters
    )
    engagement = qs.filter(id=filters["engagement"]).select_related(
        "organization_unit", "framework_version__framework", "lead_auditor"
    ).first()
    if not engagement:
        raise ValidationError({"engagement": "Audit engagement not found in your report scope."})
    rows = []
    for wp in Workpaper.objects.filter(
        engagement=engagement, deleted_at__isnull=True
    ).select_related("requirement", "control_implementation__control", "tester", "reviewed_by"):
        findings = Finding.objects.filter(
            tenant=tenant, audit_workpaper=wp, deleted_at__isnull=True
        ).count()
        rows.append([
            wp.title,
            wp.requirement.code if wp.requirement_id else "",
            wp.control_implementation.control.code if wp.control_implementation_id else "",
            wp.get_result_display(),
            wp.get_status_display(),
            _display_user(wp.tester),
            wp.conclusion,
            str(findings),
        ])
    return {
        "title": f"Audit Report — {engagement.title}",
        "subtitle": engagement.scope,
        "columns": ["Workpaper", "Requirement", "Control", "Result", "Status", "Tester", "Conclusion", "Findings"],
        "rows": rows,
        "summary": {
            "engagement_id": str(engagement.id),
            "status": engagement.status,
            "lead_auditor": _display_user(engagement.lead_auditor),
            "conclusion": engagement.conclusion,
            "row_count": len(rows),
        },
    }


def build_findings(tenant, filters):
    qs = _scope_queryset(
        Finding.objects.filter(tenant=tenant, deleted_at__isnull=True), filters
    ).select_related("owner", "organization_unit")
    if filters["status"]:
        qs = qs.filter(status=filters["status"])
    if filters["severity"]:
        qs = qs.filter(severity=filters["severity"])
    if filters["from_date"]:
        qs = qs.filter(due_date__gte=filters["from_date"])
    if filters["to_date"]:
        qs = qs.filter(due_date__lte=filters["to_date"])
    rows = [[
        x.title, x.get_finding_type_display(), x.get_severity_display(),
        x.get_status_display(), _display_user(x.owner),
        x.due_date.isoformat() if x.due_date else "", x.root_cause
    ] for x in qs.order_by("due_date", "-created_at")]
    return {
        "title": "Findings Register",
        "columns": ["Title", "Type", "Severity", "Status", "Owner", "Due date", "Root cause"],
        "rows": rows,
        "summary": {"row_count": len(rows)},
    }


def build_actions(tenant, filters):
    qs = _scope_queryset(
        Action.objects.filter(tenant=tenant, deleted_at__isnull=True), filters
    ).select_related("owner", "reviewer", "organization_unit")
    if filters["status"]:
        qs = qs.filter(status=filters["status"])
    if filters["priority"]:
        qs = qs.filter(priority=filters["priority"])
    if filters["from_date"]:
        qs = qs.filter(due_date__gte=filters["from_date"])
    if filters["to_date"]:
        qs = qs.filter(due_date__lte=filters["to_date"])
    rows = [[
        x.title, x.get_priority_display(), x.get_status_display(), _display_user(x.owner),
        _display_user(x.reviewer), x.due_date.isoformat() if x.due_date else "",
        str(x.progress), x.source_type
    ] for x in qs.order_by("due_date", "-priority", "title")]
    return {
        "title": "Actions Register",
        "columns": ["Title", "Priority", "Status", "Owner", "Reviewer", "Due date", "Progress", "Source"],
        "rows": rows,
        "summary": {"row_count": len(rows)},
    }


BUILDERS = {
    "soa": build_soa,
    "rtp": build_rtp,
    "audit": build_audit,
    "findings": build_findings,
    "actions": build_actions,
}


def table_html(report):
    esc = __import__("html").escape
    headers = "".join(f"<th>{esc(str(c))}</th>" for c in report["columns"])
    body = "".join(
        "<tr>" + "".join(f"<td>{esc(str(v or ''))}</td>" for v in row) + "</tr>"
        for row in report["rows"]
    )
    return f"""<!doctype html><html dir="rtl"><meta charset="utf-8"><body>
<h1>{esc(report["title"])}</h1>
<p>{esc(str(report.get("subtitle") or ""))}</p>
<table><thead><tr>{headers}</tr></thead><tbody>{body}</tbody></table>
</body></html>"""


def export_csv(report):
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(report["columns"])
    writer.writerows(report["rows"])
    return stream.getvalue().encode("utf-8-sig")


def export_xlsx(report):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(report["columns"])
    for row in report["rows"]:
        ws.append([str(v) if v is not None else "" for v in row])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def export_docx(report):
    from docx import Document
    doc = Document()
    doc.add_heading(report["title"], level=1)
    if report.get("subtitle"):
        doc.add_paragraph(str(report["subtitle"]))
    table = doc.add_table(rows=1, cols=len(report["columns"]))
    for index, value in enumerate(report["columns"]):
        table.rows[0].cells[index].text = str(value)
    for row in report["rows"]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value or "")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
