from io import BytesIO
from django.http import HttpResponse
from docx import Document as Docx

def soa_docx(assessment):
    from apps.controls.models import ControlRequirement,ControlImplementation
    doc=Docx();doc.add_heading(f"Statement of Applicability — {assessment.title}",0);table=doc.add_table(rows=1,cols=6);h=table.rows[0].cells
    for i,x in enumerate(['Control','Requirement','Applicable','Implementation','Effectiveness','Evidence']): h[i].text=x
    maps=ControlRequirement.objects.filter(requirement__framework_version=assessment.framework_version,deleted_at__isnull=True).select_related('control','requirement')
    for m in maps:
        impl=ControlImplementation.objects.filter(tenant=assessment.tenant,control=m.control,organization_unit=assessment.organization_unit,deleted_at__isnull=True).first() if assessment.organization_unit_id else None
        cells=table.add_row().cells;cells[0].text=f"{m.control.code} {m.control.title}";cells[1].text=m.requirement.code;cells[2].text='Yes';cells[3].text=impl.implementation_status if impl else 'not_implemented';cells[4].text=impl.effectiveness if impl else 'not_assessed';cells[5].text=''
    b=BytesIO();doc.save(b);resp=HttpResponse(b.getvalue(),content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document');resp['Content-Disposition']=f'attachment; filename="soa-{assessment.id}.docx"';return resp

def document_docx(version):
    doc=Docx();doc.add_heading(version.document.title,0);doc.add_paragraph(version.content or '');b=BytesIO();doc.save(b);resp=HttpResponse(b.getvalue(),content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document');resp['Content-Disposition']=f'attachment; filename="{version.document.code}-{version.version}.docx"';return resp
