from io import BytesIO
from django.http import HttpResponse
from rest_framework.response import Response
from .services import latest_evaluations
from apps.actions.models import Action

def rtp_payload(risk):
    evaluations=latest_evaluations(risk)
    treatments=[]
    for t in risk.treatments.filter(deleted_at__isnull=True).select_related("owner"):
        actions=Action.objects.filter(tenant=risk.tenant,source_type="risk_treatment",source_id=t.id,deleted_at__isnull=True).select_related("owner")
        treatments.append({"id":str(t.id),"strategy":t.strategy,"description":t.description,"owner":t.owner.get_full_name() or t.owner.get_username(),"target_date":t.target_date.isoformat() if t.target_date else None,"target_score":str(t.target_score) if t.target_score is not None else None,"status":t.status,"actions":[{"title":a.title,"owner":a.owner.get_full_name() or a.owner.get_username(),"due_date":a.due_date.isoformat() if a.due_date else None,"progress":a.progress,"status":a.status} for a in actions]})
    controls=[{"code":link.control_implementation.control.code,"title":link.control_implementation.control.title,"relationship_type":link.relationship_type,"effectiveness":link.control_implementation.effectiveness} for link in risk.control_links.filter(deleted_at__isnull=True).select_related("control_implementation__control")]
    return {"risk":{"id":str(risk.id),"code":risk.code,"title":risk.title,"scenario":risk.scenario,"owner":risk.owner.get_full_name() or risk.owner.get_username(),"status":risk.status},"evaluations":{k:{"score":str(v.score),"level":v.level,"likelihood":str(v.likelihood),"impact":str(v.impact),"evaluated_at":v.evaluated_at.isoformat()} for k,v in evaluations.items()},"controls":controls,"treatments":treatments}

def rtp_docx_response(risk):
    from docx import Document
    data=rtp_payload(risk);doc=Document();doc.add_heading('Risk Treatment Plan',0);doc.add_heading(f"{risk.code} — {risk.title}",level=1);doc.add_paragraph(f"Owner: {data['risk']['owner']}");doc.add_paragraph(f"Scenario: {risk.scenario or '-'}")
    doc.add_heading('Risk evaluations',level=2);table=doc.add_table(rows=1,cols=5);hdr=table.rows[0].cells
    for i,v in enumerate(['Type','Likelihood','Impact','Score','Level']): hdr[i].text=v
    for typ,row in data['evaluations'].items(): cells=table.add_row().cells;cells[0].text=typ;cells[1].text=row['likelihood'];cells[2].text=row['impact'];cells[3].text=row['score'];cells[4].text=row['level']
    doc.add_heading('Treatment plan',level=2)
    for t in data['treatments']:
        doc.add_heading(t['strategy'].title(),level=3);doc.add_paragraph(t['description'] or '-');doc.add_paragraph(f"Owner: {t['owner']} | Target: {t['target_date'] or '-'} | Status: {t['status']}")
        for a in t['actions']: doc.add_paragraph(f"• {a['title']} — {a['owner']} — {a['progress']}% — due {a['due_date'] or '-'}")
    stream=BytesIO();doc.save(stream);stream.seek(0);response=HttpResponse(stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document');response['Content-Disposition']=f'attachment; filename="RTP-{risk.code}.docx"';return response
