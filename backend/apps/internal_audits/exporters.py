from io import BytesIO
from django.http import HttpResponse
from docx import Document

def audit_report_docx(engagement):
    doc=Document(); doc.add_heading(engagement.title,0); doc.add_paragraph(f"Type: {engagement.get_audit_type_display()}"); doc.add_paragraph(f"Scope: {engagement.scope}"); doc.add_paragraph(f"Objective: {engagement.objective}")
    doc.add_heading('Workpapers',level=1)
    for wp in engagement.workpapers.filter(deleted_at__isnull=True).order_by('created_at'):
        doc.add_heading(wp.title,level=2); doc.add_paragraph(f"Result: {wp.get_result_display()}"); doc.add_paragraph(wp.conclusion or '')
    doc.add_heading('Findings',level=1)
    for wp in engagement.workpapers.filter(deleted_at__isnull=True):
        for f in wp.findings.filter(deleted_at__isnull=True):
            doc.add_heading(f.title,level=2); doc.add_paragraph(f"Severity: {f.get_severity_display()}"); doc.add_paragraph(f.description or '')
    if engagement.conclusion: doc.add_heading('Conclusion',level=1); doc.add_paragraph(engagement.conclusion)
    buf=BytesIO(); doc.save(buf); buf.seek(0); resp=HttpResponse(buf.getvalue(),content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'); resp['Content-Disposition']=f'attachment; filename="audit-{engagement.id}.docx"'; return resp
