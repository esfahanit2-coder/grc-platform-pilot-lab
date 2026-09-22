from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from apps.assets.models import Asset
from apps.audit.models import AuditEvent
from apps.identity.models import UserRoleScope
from apps.identity.services import bootstrap_tenant_rbac
from apps.organizations.models import OrganizationUnit
from apps.tenancy.models import Tenant, TenantMembership


class DataExchangeTests(APITestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(username="exchange-admin", password="StrongPassword123!")
        self.viewer = User.objects.create_user(username="exchange-viewer", password="StrongPassword123!")
        self.owner = User.objects.create_user(username="asset-owner", password="StrongPassword123!")
        self.tenant = Tenant.objects.create(name="Exchange Tenant", code="exchange")
        for user, role in ((self.admin, "admin"), (self.viewer, "viewer"), (self.owner, "member")):
            TenantMembership.objects.create(tenant=self.tenant, user=user, role_code=role)
        roles = bootstrap_tenant_rbac(self.tenant, admin_user=self.admin)
        self.hq = OrganizationUnit.objects.create(
            tenant=self.tenant, unit_type=OrganizationUnit.UnitType.COMPANY, code="HQ", name="HQ"
        )
        self.ops = OrganizationUnit.objects.create(
            tenant=self.tenant, unit_type=OrganizationUnit.UnitType.BUSINESS_UNIT, code="OPS", name="OPS"
        )
        UserRoleScope.objects.create(
            user=self.viewer, tenant=self.tenant, role=roles["viewer"], organization_unit=self.hq
        )
        Asset.objects.create(
            tenant=self.tenant, organization_unit=self.hq, asset_type=Asset.AssetType.HARDWARE,
            code="A-HQ", title="HQ Asset", owner=self.owner
        )
        Asset.objects.create(
            tenant=self.tenant, organization_unit=self.ops, asset_type=Asset.AssetType.HARDWARE,
            code="A-OPS", title="OPS Asset", owner=self.owner
        )

    def auth(self, user):
        token = str(RefreshToken.for_user(user).access_token)
        self.client.credentials(
            HTTP_AUTHORIZATION=f"Bearer {token}", HTTP_X_TENANT_ID=str(self.tenant.id)
        )

    def csv_file(self, name, text):
        stream = BytesIO(text.encode("utf-8"))
        stream.name = name
        return stream

    def xlsx_file(self, name, headers, row):
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        ws.append(row)
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        stream.name = name
        return stream

    def test_user_template_never_contains_password_or_role_columns(self):
        self.auth(self.admin)
        response = self.client.get("/api/v1/data-exchange/template/?dataset=users&format=csv")
        self.assertEqual(response.status_code, 200)
        body = b"".join(response.streaming_content) if getattr(response, "streaming", False) else response.content
        header = body.decode("utf-8-sig").splitlines()[0]
        self.assertNotIn("password", header)
        self.assertNotIn("role", header)

    def test_dry_run_does_not_mutate_and_commit_requires_same_file(self):
        self.auth(self.admin)
        text = (
            "code,title,asset_type,organization_code,owner_username,custodian_username,description,confidentiality,integrity,availability,criticality,status\n"
            "A-NEW,New Asset,hardware,HQ,asset-owner,,Imported,3,3,3,3,active\n"
        )
        response = self.client.post(
            "/api/v1/data-exchange/validate/",
            {"dataset": "assets", "file": self.csv_file("assets.csv", text)},
            format="multipart",
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["error_count"], 0)
        self.assertFalse(Asset.objects.filter(tenant=self.tenant, code="A-NEW").exists())

        changed = text.replace("New Asset", "Changed Asset")
        rejected = self.client.post(
            "/api/v1/data-exchange/commit/",
            {
                "dataset": "assets",
                "validation_token": body["validation_token"],
                "file": self.csv_file("assets.csv", changed),
            },
            format="multipart",
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertFalse(Asset.objects.filter(tenant=self.tenant, code="A-NEW").exists())

        committed = self.client.post(
            "/api/v1/data-exchange/commit/",
            {
                "dataset": "assets",
                "validation_token": body["validation_token"],
                "file": self.csv_file("assets.csv", text),
            },
            format="multipart",
        )
        self.assertEqual(committed.status_code, 200)
        self.assertTrue(Asset.objects.filter(tenant=self.tenant, code="A-NEW").exists())
        self.assertTrue(AuditEvent.objects.filter(tenant=self.tenant, action="data_exchange.import").exists())

    def test_scoped_export_excludes_other_unit(self):
        self.auth(self.viewer)
        response = self.client.get("/api/v1/data-exchange/export/?dataset=assets&format=csv")
        self.assertEqual(response.status_code, 200)
        text = response.content.decode("utf-8-sig")
        self.assertIn("A-HQ", text)
        self.assertNotIn("A-OPS", text)

    def test_password_column_is_rejected(self):
        self.auth(self.admin)
        text = "username,first_name,last_name,email,membership_active,password\nnew-user,New,User,new@example.com,true,Secret\n"
        response = self.client.post(
            "/api/v1/data-exchange/validate/",
            {"dataset": "users", "file": self.csv_file("users.csv", text)},
            format="multipart",
        )
        self.assertEqual(response.status_code, 400)

    def test_xlsx_template_is_valid_workbook(self):
        self.auth(self.admin)
        response = self.client.get("/api/v1/data-exchange/template/?dataset=assets&format=xlsx")
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content), read_only=True)
        headers = [cell.value for cell in next(wb.active.iter_rows())]
        self.assertEqual(headers[0], "code")

    def test_export_does_not_leak_other_tenant(self):
        User = get_user_model()
        other_tenant = Tenant.objects.create(name="Other", code="exchange-other")
        other_user = User.objects.create_user(username="other-owner", password="StrongPassword123!")
        TenantMembership.objects.create(tenant=other_tenant, user=other_user, role_code="admin")
        bootstrap_tenant_rbac(other_tenant, admin_user=other_user)
        other_unit = OrganizationUnit.objects.create(
            tenant=other_tenant, unit_type=OrganizationUnit.UnitType.COMPANY, code="OTHER", name="Other"
        )
        Asset.objects.create(
            tenant=other_tenant, organization_unit=other_unit, asset_type=Asset.AssetType.HARDWARE,
            code="SECRET-OTHER", title="Other Asset", owner=other_user
        )
        self.auth(self.admin)
        response = self.client.get("/api/v1/data-exchange/export/?dataset=assets&format=csv")
        self.assertNotIn("SECRET-OTHER", response.content.decode("utf-8-sig"))

    def test_new_user_import_creates_unusable_password_without_roles(self):
        self.auth(self.admin)
        text = "username,first_name,last_name,email,membership_active\nimported-user,Imported,User,imported@example.com,true\n"
        validated = self.client.post(
            "/api/v1/data-exchange/validate/",
            {"dataset": "users", "file": self.csv_file("users.csv", text)},
            format="multipart",
        )
        self.assertEqual(validated.status_code, 200)
        self.assertEqual(validated.json()["error_count"], 0)
        committed = self.client.post(
            "/api/v1/data-exchange/commit/",
            {
                "dataset": "users",
                "validation_token": validated.json()["validation_token"],
                "file": self.csv_file("users.csv", text),
            },
            format="multipart",
        )
        self.assertEqual(committed.status_code, 200)
        User = get_user_model()
        user = User.objects.get(username="imported-user")
        self.assertFalse(user.has_usable_password())
        self.assertTrue(TenantMembership.objects.filter(tenant=self.tenant, user=user, is_active=True).exists())
        self.assertFalse(UserRoleScope.objects.filter(tenant=self.tenant, user=user, is_active=True).exists())

    def test_xlsx_asset_import_dry_run_and_commit(self):
        self.auth(self.admin)
        headers = [
            "code", "title", "asset_type", "organization_code", "owner_username",
            "custodian_username", "description", "confidentiality", "integrity",
            "availability", "criticality", "status",
        ]
        row = ["A-XLSX", "XLSX Asset", "hardware", "HQ", "asset-owner", "", "", 3, 3, 3, 3, "active"]
        source = self.xlsx_file("assets.xlsx", headers, row)
        raw = source.getvalue()

        validated_file = BytesIO(raw)
        validated_file.name = "assets.xlsx"
        validated = self.client.post(
            "/api/v1/data-exchange/validate/",
            {"dataset": "assets", "file": validated_file},
            format="multipart",
        )
        self.assertEqual(validated.status_code, 200)
        self.assertEqual(validated.json()["error_count"], 0)
        self.assertFalse(Asset.objects.filter(tenant=self.tenant, code="A-XLSX").exists())

        committed_file = BytesIO(raw)
        committed_file.name = "assets.xlsx"
        committed = self.client.post(
            "/api/v1/data-exchange/commit/",
            {
                "dataset": "assets",
                "validation_token": validated.json()["validation_token"],
                "file": committed_file,
            },
            format="multipart",
        )
        self.assertEqual(committed.status_code, 200)
        self.assertTrue(Asset.objects.filter(tenant=self.tenant, code="A-XLSX").exists())
